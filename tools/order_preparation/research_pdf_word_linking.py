"""Read-only, bounded PDF-to-Word OCR research pilot.

The generated workbook is a local research artifact.  It contains only
redacted excerpts unless --include-full-text is explicitly supplied.
"""
from __future__ import annotations

import argparse
import hashlib
import re
import shutil
import subprocess
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader

from prepare_personnel_orders import action_type, clean, extract_requisites, safe_cell

NAME = re.compile(r"\b[\u0400-\u04ff][\u0400-\u04ff-]+\s+[\u0400-\u04ff][\u0400-\u04ff-]+(?:\s+[\u0400-\u04ff][\u0400-\u04ff-]+)?\b")
ORDER_START = re.compile(r"\b(?:\u041f\u0420\u0418\u041a\u0410\u0417|\u0411\u04b0\u0419\u0420\u042b\u049a)\b", re.I)


def file_sha(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for part in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(part)
    return digest.hexdigest()


def text_sha(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def redacted(value: str, limit: int = 220) -> str:
    return NAME.sub("[FIO]", clean(value))[:limit]


def ocr_capabilities() -> tuple[str | None, list[str], str]:
    executable = shutil.which("tesseract") or (r"C:\Program Files\Tesseract-OCR\tesseract.exe" if Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe").is_file() else None)
    if not executable:
        return None, [], "not installed"
    version = subprocess.run([executable, "--version"], capture_output=True, text=True, encoding="utf-8", errors="replace").stdout.splitlines()[0]
    listed = subprocess.run([executable, "--list-langs"], capture_output=True, text=True, encoding="utf-8", errors="replace").stdout.splitlines()
    return executable, [item.strip() for item in listed[1:] if item.strip()], version


def render(page: Any, dpi: int, variant: str):
    from PIL import Image, ImageEnhance, ImageOps
    pixmap = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
    image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
    gray = ImageOps.grayscale(image)
    if variant == "contrast": return ImageEnhance.Contrast(gray).enhance(2.0)
    if variant == "binary": return gray.point(lambda value: 255 if value > 160 else 0)
    return gray


def run_ocr(image: Any, executable: str | None, language: str, psm: int) -> tuple[str, float]:
    if not executable: return "", 0.0
    import pytesseract
    pytesseract.pytesseract.tesseract_cmd = executable
    try: values = pytesseract.image_to_data(image, lang=language, config=f"--psm {psm}", output_type=pytesseract.Output.DICT)
    except Exception: return "", 0.0
    words = [item for item in values["text"] if item.strip()]
    confidence = [float(item) for item in values["conf"] if str(item) not in {"-1", ""}]
    return " ".join(words), (sum(confidence) / len(confidence) if confidence else 0.0)


def chosen_pages(path: Path) -> list[int]:
    count = len(PdfReader(str(path)).pages)
    return sorted({0, max(0, (count - 1) // 2), max(0, count - 1)})


def word_info(path: Path) -> dict[str, Any]:
    try:
        document = Document(path); paragraphs = [item.text for item in document.paragraphs if item.text.strip()]
        cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells if cell.text.strip()]
        text = "\n".join([*paragraphs, *cells]); _, number, _, date, _, _ = extract_requisites(text)
        return {"paragraphs": len(paragraphs), "tables": len(document.tables), "sections": len(document.sections), "images": len(document.inline_shapes), "number": number, "date": date, "action": action_type(text), "text": text, "warning": ""}
    except Exception:
        return {"paragraphs": 0, "tables": 0, "sections": 0, "images": 0, "number": "", "date": "", "action": "", "text": "", "warning": "DOCX_READ_FAILED"}


def main() -> int:
    parser = argparse.ArgumentParser(); parser.add_argument("root", type=Path); parser.add_argument("--output", type=Path, required=True); parser.add_argument("--dpi", type=int, default=300); parser.add_argument("--include-full-text", action="store_true", help="Add a local sheet containing personal data.")
    args = parser.parse_args(); files = sorted(path for path in args.root.rglob("*") if path.is_file()); pdfs = [path for path in files if path.suffix.lower() == ".pdf"]; docs = [path for path in files if path.suffix.lower() == ".docx"]
    executable, languages, version = ocr_capabilities(); language = "rus+eng" if {"rus", "eng"}.issubset(languages) else "rus" if "rus" in languages else "eng"
    pdf_rows: list[list[Any]] = []; orders: list[list[Any]] = []; experiments: list[list[Any]] = []; review: list[list[Any]] = []; full_text: list[list[Any]] = []
    for pdf in pdfs:
        reader = PdfReader(str(pdf)); text_layer = "".join(page.extract_text() or "" for page in reader.pages); rel = pdf.relative_to(args.root).as_posix()
        pdf_rows.append([rel, pdf.stat().st_size, file_sha(pdf), len(reader.pages), bool(clean(text_layer)), "pages 1/middle/last"])
        document = fitz.open(pdf)
        for page_index in chosen_pages(pdf):
            page = document[page_index]; best = ("", 0.0, "", 0, 0)
            for variant, dpi, psm in [("gray", args.dpi, 6), ("contrast", args.dpi, 6), ("binary", args.dpi, 6), ("gray", 400, 4)]:
                text, confidence = run_ocr(render(page, dpi, variant), executable, language, psm); _, number, _, date, _, _ = extract_requisites(text); signals = sum(bool(item) for item in (number, date, ORDER_START.search(text)))
                experiments.append([rel, page_index + 1, variant, dpi, psm, language, round(confidence, 1), len(text), signals, text_sha(text)])
                old_text, old_confidence, _, _, _ = best; _, old_number, _, old_date, _, _ = extract_requisites(old_text); old_signals = sum(bool(item) for item in (old_number, old_date, ORDER_START.search(old_text)))
                if (signals, confidence, len(text)) > (old_signals, old_confidence, len(old_text)): best = (text, confidence, variant, dpi, psm)
            text, confidence, variant, dpi, psm = best; _, number, _, date, _, _ = extract_requisites(text); start = bool(ORDER_START.search(text) or (number and date)); row_review = "CANDIDATE_ONLY" if number and date and confidence >= 70 else "REVIEW_REQUIRED"
            orders.append([rel, page_index + 1, number, date, action_type(text), round(confidence, 1), variant, dpi, psm, start, redacted(text), text_sha(text), row_review])
            if args.include_full_text: full_text.append([rel, page_index + 1, language, variant, dpi, psm, text_sha(text), text])
            if row_review == "REVIEW_REQUIRED": review.append(["PDF", rel, page_index + 1, "OCR_REQUISITES_INCOMPLETE", redacted(text)])
    word_rows: list[list[Any]] = []; word_orders: list[list[Any]] = []
    for document_path in docs:
        info = word_info(document_path); rel = document_path.relative_to(args.root).as_posix()
        word_rows.append([rel, document_path.stat().st_size, file_sha(document_path), info["paragraphs"], info["tables"], info["sections"], info["images"], info["warning"]]); word_orders.append([rel, "document", info["number"], info["date"], info["action"], redacted(info["text"]), text_sha(info["text"]), "CANDIDATE_ONLY"])
    candidates: list[list[Any]] = []
    for order in orders:
        matched = [word for word in word_orders if order[2] and order[3] and word[2] == order[2] and word[3] == order[3]]
        status = "STRONG_CANDIDATE" if len(matched) == 1 and order[5] >= 70 else "REVIEW_REQUIRED" if len(matched) > 1 or matched else "OCR_FAILED" if not order[2] or not order[3] else "NO_CANDIDATE"
        candidates.append([order[0], order[1], matched[0][0] if len(matched) == 1 else "", len(matched), order[2], order[3], order[4], status, "number+date exact candidate; never confirmed"])
    tabs = {"PDFDocuments": (["path", "size", "sha256", "pages", "text_layer", "pilot"], pdf_rows), "PDFOrders": (["pdf", "page", "number", "date", "action", "ocr_confidence", "variant", "dpi", "psm", "order_start_candidate", "redacted_excerpt", "ocr_sha256", "review"], orders), "WordDocuments": (["path", "size", "sha256", "paragraphs", "tables", "sections", "images", "warning"], word_rows), "WordOrders": (["word", "block", "number", "date", "action", "redacted_excerpt", "text_sha256", "review"], word_orders), "MatchCandidates": (["pdf", "page", "word", "candidate_count", "number", "date", "action", "status", "explanation"], candidates), "ManualReview": (["source_type", "path", "page", "reason", "redacted_excerpt"], review), "OCRExperiments": (["pdf", "page", "variant", "dpi", "psm", "languages", "mean_confidence", "chars", "feature_count", "text_sha256"], experiments)}
    if args.include_full_text: tabs["OCRFullText_LOCAL_PII"] = (["pdf", "page", "languages", "variant", "dpi", "psm", "text_sha256", "full_ocr_text"], full_text)
    book = Workbook(); book.remove(book.active)
    for title, (header, rows) in tabs.items():
        sheet = book.create_sheet(title); sheet.append(header)
        for row in rows: sheet.append([safe_cell(value) for value in row])
    metrics = book.create_sheet("Metrics"); metrics.append(["metric", "value"])
    for pair in [("pilot_pages", len(orders)), ("pdf_files", len(pdfs)), ("word_docx", len(docs)), ("strong_candidates", sum(row[7] == "STRONG_CANDIDATE" for row in candidates)), ("review_required", sum(row[7] == "REVIEW_REQUIRED" for row in candidates)), ("no_candidate_or_ocr_failed", sum(row[7] in {"NO_CANDIDATE", "OCR_FAILED"} for row in candidates))]: metrics.append(pair)
    metadata = book.create_sheet("Metadata"); metadata.append(["key", "value"])
    for pair in [("tesseract", version), ("languages", ",".join(languages)), ("selected_languages", language), ("source_root", str(args.root)), ("privacy", "OCRFullText_LOCAL_PII contains personal data; do not distribute" if args.include_full_text else "redacted excerpts plus hashes; no full OCR text")]: metadata.append(pair)
    args.output.parent.mkdir(parents=True, exist_ok=True); book.save(args.output); print(f"pdf={len(pdfs)} docs={len(docs)} pilot_pages={len(orders)} output={args.output}")
    return 0


if __name__ == "__main__": raise SystemExit(main())
