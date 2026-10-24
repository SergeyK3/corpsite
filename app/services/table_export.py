"""Render already authorized table values; this module never reads application data."""
from io import BytesIO
from itertools import chain
from collections.abc import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter


def build_table_workbook(headers: Sequence[str], rows: Iterable[Sequence], title: str = "Данные", widths: Sequence[int] | None = None) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = "A2"
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1] if widths else 30
    count = 0
    for values in chain((headers,), rows):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            cells.append(cell)
        sheet.append(cells)
        count += 1
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
