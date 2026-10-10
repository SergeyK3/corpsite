"""Strict, non-destructive import of the reviewed job catalog."""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

FIELDS = ("job_code", "job_nameru", "job_namekk", "job_namekk_doc", "legacy_position_ids")


def read_catalog(path):
    if isinstance(path, (str, Path)) and Path(path).suffix.lower() == ".json":
        package = json.loads(Path(path).read_text(encoding="utf-8"))
        if package.get("format") != "job-positions-catalog-v1":
            raise ValueError("Unexpected catalog package format")
        rows = package.get("items")
        if not isinstance(rows, list) or not rows:
            raise ValueError("Catalog package has no rows")
        for number, row in enumerate(rows, 1):
            if not isinstance(row, dict) or set(row) != set(FIELDS):
                raise ValueError(f"Unexpected catalog fields at row {number}")
            if any(not isinstance(row[key], str) or not row[key].strip() for key in FIELDS[:4]):
                raise ValueError(f"Empty or invalid catalog text at row {number}")
            if not re.fullmatch(r"[A-Z][A-Z0-9_]*", row["job_code"]):
                raise ValueError(f"Invalid job code at row {number}")
            ids = row["legacy_position_ids"]
            if not isinstance(ids, list) or any(type(value) is not int or value <= 0 for value in ids):
                raise ValueError(f"Invalid legacy IDs at row {number}")
        return rows
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook["catl_job"]
        if sheet.tables["CatlJobPositionsV2"].ref != "A1:E89":
            raise ValueError("Expected CatlJobPositionsV2 at A1:E89")
        cells = list(sheet.iter_rows(min_row=1, max_row=89, max_col=5))
        if tuple(c.value for c in cells[0]) != FIELDS:
            raise ValueError("Unexpected catalog columns")
        rows = []
        for number, cells_row in enumerate(cells[1:], 2):
            if any(c.data_type == "f" for c in cells_row):
                raise ValueError(f"Formula in source row {number}")
            values = [str(c.value).strip() if c.value is not None else "" for c in cells_row]
            if not all(values[:4]) or not re.fullmatch(r"[A-Z][A-Z0-9_]*", values[0]):
                raise ValueError(f"Invalid catalog row {number}")
            raw_ids = values[4]
            if raw_ids and not re.fullmatch(r"[1-9][0-9]*(\s*,\s*[1-9][0-9]*)*", raw_ids):
                raise ValueError(f"Invalid legacy IDs at row {number}")
            row = dict(zip(FIELDS, values))
            row["legacy_position_ids"] = [int(v.strip()) for v in raw_ids.split(",")] if raw_ids else []
            rows.append(row)
        return rows
    finally:
        workbook.close()


def plan_catalog(rows, positions=None, catalog=(), links=()):
    existing = None if positions is None else {int(p["position_id"]): dict(p) for p in positions}
    codes = {r["job_code"]: dict(r) for r in catalog}
    linked = {int(r["position_id"]): r["job_code"] for r in links}
    seen_codes, seen_ids, seen_names = set(), {}, {}
    conflicts, mappings, merges, missing = [], [], [], []
    for row in rows:
        code = row["job_code"]
        if code in seen_codes:
            conflicts.append(f"Duplicate job_code: {code}")
        seen_codes.add(code)
        name = row["job_nameru"].casefold()
        if name in seen_names:
            conflicts.append(f"Duplicate RU name: {seen_names[name]} / {code}")
        seen_names[name] = code
        if code in codes and any(codes[code].get(k) != row[k] for k in FIELDS[1:4]):
            conflicts.append(f"Existing catalog values differ (preserved): {code}")
        matched = []
        for pid in row["legacy_position_ids"]:
            if pid in seen_ids:
                conflicts.append(f"Legacy ID {pid} repeated: {seen_ids[pid]} / {code}")
            seen_ids[pid] = code
            if pid in linked and linked[pid] != code:
                conflicts.append(f"Legacy ID {pid} already mapped to {linked[pid]}, requested {code}")
            if existing is not None and pid not in existing:
                missing.append(pid)
            matched.append({"position_id": pid, "existing_name": existing[pid].get("name") if existing is not None and pid in existing else None,
                            "mapping_action": "keep" if linked.get(pid) == code else "conflict" if pid in linked else "create",
                            "status": "unchecked" if existing is None else "found" if pid in existing else "missing"})
        mappings.append({**row, "action": "keep" if code in codes else "create", "matches": matched})
        if len(matched) > 1:
            merges.append({"job_code": code, "legacy_position_ids": row["legacy_position_ids"]})
    new_jobs = sum(row["action"] == "create" for row in mappings)
    new_links = sum(match["mapping_action"] == "create" for row in mappings for match in row["matches"])
    return {"source": "catl_job!A1:E89 / CatlJobPositionsV2", "rows": len(rows),
            "planned_changes": {"catalog_inserts": new_jobs, "mapping_inserts": new_links, "total": new_jobs + new_links, "position_updates": 0, "assignment_updates": 0, "deletes": 0},
            "database_checked": existing is not None, "mappings": mappings, "merges": merges,
            "missing_ids": sorted(set(missing)) if existing is not None else None,
            "unmapped_existing_ids": sorted(set(existing) - set(seen_ids)) if existing is not None else None,
            "conflicts": conflicts, "can_apply": existing is not None and not conflicts and not missing}


def database_plan(conn, rows):
    positions = conn.execute(text("SELECT position_id, name FROM public.positions ORDER BY position_id")).mappings().all()
    present = conn.execute(text("SELECT to_regclass('public.job_positions_catalog')")).scalar()
    catalog = conn.execute(text("SELECT * FROM public.job_positions_catalog")).mappings().all() if present else []
    links = conn.execute(text("SELECT * FROM public.position_job_catalog")).mappings().all() if present else []
    plan = plan_catalog(rows, positions, catalog, links)
    plan["schema_ready"] = bool(present)
    plan["can_apply"] = plan["can_apply"] and bool(present)
    return plan


def apply_catalog(conn, rows):
    # Serialize imports and UI edits; revalidate inside the same transaction.
    conn.execute(text("LOCK TABLE public.job_positions_catalog, public.position_job_catalog IN EXCLUSIVE MODE"))
    conn.execute(text("LOCK TABLE public.positions IN SHARE MODE"))
    plan = database_plan(conn, rows)
    if not plan["can_apply"]:
        raise ValueError("Import blocked: inspect missing IDs and conflicts in dry-run")
    for row in rows:
        conn.execute(text("""INSERT INTO public.job_positions_catalog(job_code, job_nameru, job_namekk, job_namekk_doc)
            VALUES (:job_code, :job_nameru, :job_namekk, :job_namekk_doc) ON CONFLICT (job_code) DO NOTHING"""), row)
        for pid in row["legacy_position_ids"]:
            conn.execute(text("""INSERT INTO public.position_job_catalog(position_id, job_code)
                VALUES (:position_id, :job_code) ON CONFLICT (position_id) DO NOTHING"""), {"position_id": pid, "job_code": row["job_code"]})
    return plan


def source_sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()
