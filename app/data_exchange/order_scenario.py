"""Versioned workbook scenario for externally prepared order review records.

It deliberately creates only ``personnel_order_import_review_records``.  The
canonical personnel-order command flow remains the sole producer of official
orders and employee events.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import PurePath, PureWindowsPath
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from app.data_exchange.contracts import PreviewRow, ScenarioResult
from app.data_exchange.permissions import (
    DATA_EXCHANGE_CONFIRM_APPLY,
    DATA_EXCHANGE_DRY_RUN,
    DATA_EXCHANGE_EXPORT_REFERENCE,
    DATA_EXCHANGE_UPLOAD_PREVIEW,
)


ORDER_SCHEMA_VERSION = "PERSONNEL_ORDERS_XLSX_V1"
ORDER_SCENARIO_CODE = "PERSONNEL_ORDERS_REVIEW_IMPORT"
EMPLOYEE_EXPORT_SCHEMA_VERSION = "EMPLOYEE_REFERENCE_XLSX_V1"
EMPLOYEE_EXPORT_SCENARIO_CODE = "EMPLOYEE_REFERENCE_EXPORT"
MAX_XLSX_BYTES = 25 * 1024 * 1024
_SHA_RE = re.compile(r"^[0-9a-f]{64}$")
_DANGEROUS_EXCEL_PREFIX = ("=", "+", "-", "@")
_REQUIRED_ORDER_COLUMNS = (
    "schema_version", "source_row_id", "employee_id", "employee_full_name",
    "order_number_raw", "order_number", "order_date_raw", "order_date",
    "order_kind", "action_type", "pdf_source", "pdf_sha256", "word_source", "word_sha256",
    "source_text", "recognition_confidence", "recognition_status",
    "matching_method", "review_status", "note",
)


class WorkbookValidationError(ValueError):
    """Safe structural error returned to the operator without workbook values."""


def _normal_text(value: Any) -> str:
    return " ".join(str(value or "").split()).casefold()


def _safe_relative_source(value: Any) -> bool:
    raw = str(value or "").strip()
    if not raw or raw.startswith(("/", "\\")) or ".." in raw.replace("\\", "/").split("/"):
        return False
    windows = PureWindowsPath(raw)
    posix = PurePath(raw)
    return not windows.drive and not windows.root and not posix.is_absolute()


def _date_value(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def _safe_cell_text(value: Any) -> str:
    """Escape text only when writing workbooks, never alter source evidence."""
    rendered = "" if value is None else str(value)
    return "'" + rendered if rendered.startswith(_DANGEROUS_EXCEL_PREFIX) else rendered


def _row_fingerprint(payload: dict[str, Any]) -> str:
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _open_xlsx(content: bytes):
    if len(content) == 0 or len(content) > MAX_XLSX_BYTES or not content.startswith(b"PK"):
        raise WorkbookValidationError("Unsupported or oversized workbook.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            # The compressed upload limit is not a zip-bomb protection.  Keep
            # this deliberately modest because the server only needs the
            # small, tabular interchange workbook.
            if len(archive.infolist()) > 200 or sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                raise WorkbookValidationError("Workbook archive is too large.")
            if any(name.lower().endswith("vbaproject.bin") for name in archive.namelist()):
                raise WorkbookValidationError("Macro-enabled workbooks are not accepted.")
        return load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_vba=False)
    except WorkbookValidationError:
        raise
    except Exception as exc:
        raise WorkbookValidationError("Workbook cannot be read.") from exc


def _metadata_schema(workbook) -> str:
    if "Metadata" not in workbook.sheetnames:
        raise WorkbookValidationError("Required Metadata worksheet is missing.")
    metadata = workbook["Metadata"]
    values = {
        str(key.value or "").strip(): str(value.value or "").strip()
        for key, value in metadata.iter_rows(min_row=1, max_col=2)
        if key.value is not None
    }
    return values.get("schema_version", "")


def _read_order_rows(content: bytes) -> list[tuple[int, dict[str, Any]]]:
    workbook = _open_xlsx(content)
    if _metadata_schema(workbook) != ORDER_SCHEMA_VERSION:
        raise WorkbookValidationError("Unsupported order workbook schema version.")
    if "Orders" not in workbook.sheetnames:
        raise WorkbookValidationError("Required Orders worksheet is missing.")
    sheet = workbook["Orders"]
    header_cells = next(sheet.iter_rows(min_row=1, max_row=1))
    headers = [str(cell.value or "").strip() for cell in header_cells]
    if tuple(headers[:len(_REQUIRED_ORDER_COLUMNS)]) != _REQUIRED_ORDER_COLUMNS:
        raise WorkbookValidationError("Orders worksheet columns do not match the schema.")
    results: list[tuple[int, dict[str, Any]]] = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        if all(cell.value is None for cell in cells):
            continue
        if any(cell.data_type == "f" for cell in cells):
            raise WorkbookValidationError("Workbook formulas are not accepted.")
        values = {headers[index]: cells[index].value for index in range(len(_REQUIRED_ORDER_COLUMNS))}
        results.append((row_number, values))
    if not results:
        raise WorkbookValidationError("Orders worksheet has no rows.")
    return results


def _employee_snapshot(conn, employee_ids: set[int], *, scope_unit_ids: set[int] | None) -> dict[int, dict[str, Any]]:
    if not employee_ids:
        return {}
    rows = conn.execute(text("""
        SELECT e.employee_id, e.person_id, COALESCE(p.full_name, e.full_name) AS full_name,
               p.updated_at AS person_updated_at
        FROM public.employees e
        LEFT JOIN public.persons p ON p.person_id=e.person_id
        WHERE e.employee_id = ANY(:employee_ids)
          AND (:scope_unit_ids IS NULL OR e.org_unit_id = ANY(:scope_unit_ids))
    """), {"employee_ids": sorted(employee_ids), "scope_unit_ids": sorted(scope_unit_ids) if scope_unit_ids is not None else None}).mappings().all()
    return {int(row["employee_id"]): dict(row) for row in rows}


def _target_fingerprint(snapshot: dict[int, dict[str, Any]]) -> str:
    material = [
        {"employee_id": employee_id, "person_id": row.get("person_id"), "full_name": row.get("full_name"),
         "person_updated_at": str(row.get("person_updated_at") or "")}
        for employee_id, row in sorted(snapshot.items())
    ]
    return _row_fingerprint({"employees": material})


def _valid_sha(value: Any) -> bool:
    return bool(_SHA_RE.fullmatch(str(value or "").strip().lower()))


def _business_key(payload: dict[str, Any]) -> str:
    return _row_fingerprint({
        "employee_id": payload["employee_id"], "order_kind": payload["order_kind"],
        "order_number": payload["order_number"], "order_date": payload["order_date"],
        "action_type": payload["action_type"], "pdf_sha256": payload["pdf_sha256"],
        "word_sha256": payload["word_sha256"],
    })


class PersonnelOrdersReviewScenario:
    code = ORDER_SCENARIO_CODE
    schema_version = ORDER_SCHEMA_VERSION
    label = "Кадровые и личные приказы (review-staging)"
    accepted_suffixes = (".xlsx",)
    max_bytes = MAX_XLSX_BYTES
    atomic_apply = True
    upload_permission = DATA_EXCHANGE_UPLOAD_PREVIEW
    dry_run_permission = DATA_EXCHANGE_DRY_RUN
    apply_permission = DATA_EXCHANGE_CONFIRM_APPLY
    export_permission = None

    def validate_and_preview(self, conn, content: bytes, *, scope_unit_ids: set[int] | None) -> ScenarioResult:
        raw_rows = _read_order_rows(content)
        employee_ids: set[int] = set()
        for _, values in raw_rows:
            try:
                employee_ids.add(int(values.get("employee_id")))
            except (TypeError, ValueError):
                pass
        snapshot = _employee_snapshot(conn, employee_ids, scope_unit_ids=scope_unit_ids)
        source_ids = Counter(str(values.get("source_row_id") or "").strip() for _, values in raw_rows)
        preview_rows: list[PreviewRow] = []
        for row_number, values in raw_rows:
            source_id = str(values.get("source_row_id") or "").strip()
            source_keys = {"source_row_id": source_id, "employee_id": values.get("employee_id"),
                           "order_number": values.get("order_number"), "order_date": values.get("order_date")}
            try:
                employee_id = int(values.get("employee_id"))
            except (TypeError, ValueError):
                preview_rows.append(PreviewRow(row_number, source_keys, None, "ERROR", "BLOCK", "EMPLOYEE_ID_INVALID", "Invalid employee_id."))
                continue
            order_date = _date_value(values.get("order_date"))
            if source_ids[source_id] != 1 or not source_id:
                preview_rows.append(PreviewRow(row_number, source_keys, None, "POSSIBLE_DUPLICATE", "BLOCK", "SOURCE_ROW_DUPLICATE", "Duplicate or empty source_row_id."))
                continue
            if str(values.get("schema_version") or "").strip() != ORDER_SCHEMA_VERSION:
                preview_rows.append(PreviewRow(row_number, source_keys, None, "ERROR", "BLOCK", "SCHEMA_VERSION_MISMATCH", "Row schema version differs."))
                continue
            if not str(values.get("order_number") or "").strip() or order_date is None or not str(values.get("action_type") or "").strip():
                preview_rows.append(PreviewRow(row_number, source_keys, None, "ERROR", "BLOCK", "ORDER_REQUISITES_INVALID", "Order number, date and action are required."))
                continue
            # PDF is mandatory evidence.  A Word source is optional because
            # some legitimate orders exist only as a signed scan; when present
            # it has the same path/hash proof requirements.
            if not _safe_relative_source(values.get("pdf_source")):
                preview_rows.append(PreviewRow(row_number, source_keys, None, "ERROR", "BLOCK", "SOURCE_PATH_UNSAFE", "Source reference must be a relative path."))
                continue
            word_source = str(values.get("word_source") or "").strip()
            word_sha = str(values.get("word_sha256") or "").strip()
            if bool(word_source) != bool(word_sha) or (word_source and not _safe_relative_source(word_source)):
                preview_rows.append(PreviewRow(row_number, source_keys, None, "ERROR", "BLOCK", "WORD_SOURCE_INVALID", "Word source and hash must be supplied together."))
                continue
            if not _valid_sha(values.get("pdf_sha256")) or (word_sha and not _valid_sha(word_sha)):
                preview_rows.append(PreviewRow(row_number, source_keys, None, "ERROR", "BLOCK", "SOURCE_HASH_INVALID", "Source SHA-256 is invalid."))
                continue
            employee = snapshot.get(employee_id)
            if employee is None:
                preview_rows.append(PreviewRow(row_number, source_keys, None, "NOT_FOUND", "BLOCK", "EMPLOYEE_NOT_FOUND", "Employee was not found."))
                continue
            if _normal_text(values.get("employee_full_name")) != _normal_text(employee.get("full_name")):
                preview_rows.append(PreviewRow(row_number, source_keys, {"employee_id": employee_id}, "AMBIGUOUS", "BLOCK", "EMPLOYEE_NAME_MISMATCH", "Employee control name does not match."))
                continue
            payload = {
                "employee_id": employee_id, "employee_full_name": str(employee["full_name"]),
                "order_kind": str(values.get("order_kind") or "PERSONNEL").strip().upper(),
                "order_number_raw": str(values.get("order_number_raw") or ""),
                "order_number": str(values.get("order_number")).strip(),
                "order_date_raw": str(values.get("order_date_raw") or ""),
                "order_date": order_date.isoformat(), "action_type": str(values.get("action_type")).strip().upper(),
                "pdf_source": str(values.get("pdf_source")).strip(), "pdf_sha256": str(values.get("pdf_sha256")).strip().lower(),
                "word_source": str(values.get("word_source")).strip(), "word_sha256": str(values.get("word_sha256")).strip().lower(),
                "source_text": str(values.get("source_text") or ""),
                "recognition_confidence": values.get("recognition_confidence"),
                "recognition_status": str(values.get("recognition_status") or ""),
                "matching_method": str(values.get("matching_method") or ""),
                "review_status": str(values.get("review_status") or ""), "note": str(values.get("note") or ""),
                "source_row_id": source_id,
            }
            key = _business_key(payload)
            exists = conn.execute(text("""SELECT review_record_id FROM public.personnel_order_import_review_records
                WHERE business_key=:business_key LIMIT 1"""), {"business_key": key}).scalar_one_or_none()
            if exists is not None:
                preview_rows.append(PreviewRow(row_number, source_keys, {"employee_id": employee_id, "review_record_id": int(exists)}, "UNCHANGED", "SKIP", "BUSINESS_KEY_EXISTS", "Identical review record already exists.", payload))
            else:
                preview_rows.append(PreviewRow(row_number, source_keys, {"employee_id": employee_id}, "CREATE", "CREATE_REVIEW_RECORD", None, None, payload))
        return ScenarioResult(preview_rows, _target_fingerprint(snapshot), [])

    def apply(self, conn, *, package_id: int, rows: list[PreviewRow], actor_user_id: int) -> dict[str, int]:
        created = 0
        skipped = 0
        for row in rows:
            if row.result_group == "UNCHANGED":
                skipped += 1
                continue
            if row.result_group != "CREATE" or not row.payload:
                raise WorkbookValidationError("Blocked preview rows cannot be applied.")
            payload = row.payload
            key = _business_key(payload)
            inserted = conn.execute(text("""
                INSERT INTO public.personnel_order_import_review_records(
                    package_id, source_row_number, business_key, employee_id, order_kind,
                    order_number, order_date, action_type, source_payload, created_by_user_id)
                VALUES(:package_id,:source_row_number,:business_key,:employee_id,:order_kind,
                       :order_number,:order_date,:action_type,CAST(:source_payload AS jsonb),:actor_user_id)
                ON CONFLICT (business_key) DO NOTHING
                RETURNING review_record_id
            """), {"package_id": package_id, "source_row_number": row.source_row_number,
                    "business_key": key, "employee_id": payload["employee_id"], "order_kind": payload["order_kind"],
                    "order_number": payload["order_number"], "order_date": payload["order_date"],
                    "action_type": payload["action_type"], "source_payload": json.dumps(payload, ensure_ascii=False),
                    "actor_user_id": actor_user_id}).scalar_one_or_none()
            if inserted is None:
                skipped += 1
            else:
                created += 1
        return {"created": created, "updated": 0, "skipped": skipped}


def build_employee_reference_workbook(conn, *, scope_unit_ids: set[int] | None) -> bytes:
    """Create the only allowed matching reference: technical ID plus control name."""
    rows = conn.execute(text("""
        SELECT e.employee_id, p.person_id, COALESCE(p.full_name, e.full_name) AS employee_full_name
        FROM public.employees e
        LEFT JOIN public.persons p ON p.person_id=e.person_id
        WHERE COALESCE(e.is_active, TRUE)=TRUE
          AND (:scope_unit_ids IS NULL OR e.org_unit_id = ANY(:scope_unit_ids))
        ORDER BY e.employee_id
    """), {"scope_unit_ids": sorted(scope_unit_ids) if scope_unit_ids is not None else None}).mappings().all()
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Employees")
    sheet.append(["schema_version", "employee_id", "person_id", "employee_full_name"])
    for row in rows:
        sheet.append([EMPLOYEE_EXPORT_SCHEMA_VERSION, int(row["employee_id"]), row["person_id"], _safe_cell_text(row["employee_full_name"])])
    metadata = workbook.create_sheet("Metadata")
    metadata.append(["schema_version", EMPLOYEE_EXPORT_SCHEMA_VERSION])
    metadata.append(["row_count", len(rows)])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
