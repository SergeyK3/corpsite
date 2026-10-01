"""Control-list qualification-category scenario with an exact HR-row proof."""
from __future__ import annotations

import hashlib
import io
import json
from collections import Counter
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text

from app.data_exchange.contracts import PreviewRow, ScenarioResult
from app.data_exchange.order_scenario import (
    MAX_XLSX_BYTES, WorkbookValidationError, _normal_text, _open_xlsx,
    _safe_cell_text,
)
from app.data_exchange.permissions import (
    DATA_EXCHANGE_CONFIRM_APPLY, DATA_EXCHANGE_DRY_RUN, DATA_EXCHANGE_UPLOAD_PREVIEW,
)
from app.personnel_intake.domain.additional_profile import normalize_additional_profile
from app.services.hr_import_qualification_category_service import parse_qualification_categories

CATEGORY_SCENARIO_CODE = "CONTROL_LIST_CATEGORIES_IMPORT"
CATEGORY_SCHEMA_VERSION = "CONTROL_LIST_CATEGORIES_XLSX_V1"
_COLUMNS = (
    "schema_version", "source_row_id", "employee_id", "employee_full_name",
    "source_batch_id", "source_import_row_id", "category_raw",
)


def _read(content: bytes) -> list[tuple[int, dict[str, Any]]]:
    workbook = _open_xlsx(content)
    if "Metadata" not in workbook.sheetnames or "Categories" not in workbook.sheetnames:
        raise WorkbookValidationError("Required Categories and Metadata worksheets are missing.")
    metadata = {str(a.value or "").strip(): str(b.value or "").strip()
                for a, b in workbook["Metadata"].iter_rows(min_row=1, max_col=2) if a.value is not None}
    if metadata.get("schema_version") != CATEGORY_SCHEMA_VERSION:
        raise WorkbookValidationError("Unsupported category workbook schema version.")
    sheet = workbook["Categories"]
    headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if tuple(headers[:len(_COLUMNS)]) != _COLUMNS:
        raise WorkbookValidationError("Categories worksheet columns do not match the schema.")
    result = []
    for number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        if all(cell.value is None for cell in cells):
            continue
        if any(cell.data_type == "f" for cell in cells):
            raise WorkbookValidationError("Workbook formulas are not accepted.")
        result.append((number, {headers[index]: cells[index].value for index in range(len(_COLUMNS))}))
    if not result:
        raise WorkbookValidationError("Categories worksheet has no rows.")
    return result


def _fingerprint(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()


class ControlListCategoriesScenario:
    code = CATEGORY_SCENARIO_CODE
    schema_version = CATEGORY_SCHEMA_VERSION
    label = "Категории из контрольного списка"
    accepted_suffixes = (".xlsx",)
    max_bytes = MAX_XLSX_BYTES
    atomic_apply = True
    upload_permission = DATA_EXCHANGE_UPLOAD_PREVIEW
    dry_run_permission = DATA_EXCHANGE_DRY_RUN
    apply_permission = DATA_EXCHANGE_CONFIRM_APPLY
    export_permission = None

    def validate_and_preview(self, conn, content: bytes, *, scope_unit_ids: set[int] | None) -> ScenarioResult:
        source_rows = _read(content)
        source_ids = Counter(str(values.get("source_row_id") or "").strip() for _, values in source_rows)
        output: list[PreviewRow] = []
        fingerprint_material: list[dict[str, Any]] = []
        for number, values in source_rows:
            source_id = str(values.get("source_row_id") or "").strip()
            keys = {"source_row_id": source_id, "employee_id": values.get("employee_id"), "source_batch_id": values.get("source_batch_id")}
            try:
                employee_id, batch_id, row_id = (int(values["employee_id"]), int(values["source_batch_id"]), int(values["source_import_row_id"]))
            except (KeyError, TypeError, ValueError):
                output.append(PreviewRow(number, keys, None, "ERROR", "BLOCK", "TECHNICAL_LINK_INVALID", "Technical source link is invalid.")); continue
            if not source_id or source_ids[source_id] != 1:
                output.append(PreviewRow(number, keys, None, "POSSIBLE_DUPLICATE", "BLOCK", "SOURCE_ROW_DUPLICATE", "Duplicate or empty source_row_id.")); continue
            row = conn.execute(text("""
                SELECT r.normalized_payload, r.employee_id AS source_employee_id, e.person_id,
                       COALESCE(p.full_name,e.full_name) AS full_name, e.org_unit_id,
                       m.additional_profile, m.updated_at AS metadata_updated_at
                FROM public.hr_import_rows r
                JOIN public.employees e ON e.employee_id=r.employee_id
                LEFT JOIN public.persons p ON p.person_id=e.person_id
                LEFT JOIN public.personnel_record_metadata m ON m.person_id=e.person_id
                WHERE r.batch_id=:batch_id AND r.row_id=:row_id
            """), {"batch_id": batch_id, "row_id": row_id}).mappings().one_or_none()
            if row is None or int(row["source_employee_id"] or 0) != employee_id:
                output.append(PreviewRow(number, keys, None, "NOT_FOUND", "BLOCK", "SOURCE_ROW_NOT_EXACTLY_BOUND", "Control-list source row is not exactly bound to employee.")); continue
            if scope_unit_ids is not None and int(row["org_unit_id"] or 0) not in scope_unit_ids:
                output.append(PreviewRow(number, keys, None, "NOT_FOUND", "BLOCK", "OUTSIDE_ORGANIZATIONAL_SCOPE", "Employee is outside organisational scope.")); continue
            if str(values.get("schema_version") or "").strip() != CATEGORY_SCHEMA_VERSION or _normal_text(values.get("employee_full_name")) != _normal_text(row["full_name"]):
                output.append(PreviewRow(number, keys, {"employee_id": employee_id}, "AMBIGUOUS", "BLOCK", "EMPLOYEE_CONTROL_MISMATCH", "Employee control values do not match.")); continue
            payload = row["normalized_payload"] or {}
            authoritative_raw = str(payload.get("certification_raw") or "").strip()
            if not authoritative_raw or str(values.get("category_raw") or "").strip() != authoritative_raw:
                output.append(PreviewRow(number, keys, {"employee_id": employee_id}, "ERROR", "BLOCK", "SOURCE_VALUE_MISMATCH", "Category value does not match the exact control-list row.")); continue
            profile = normalize_additional_profile(row["additional_profile"] or {})
            existing = profile.get("qualification_categories") or []
            if any((item.get("provenance") or {}).get("override_origin") == "HR_CORRECTION" for item in existing if isinstance(item, dict)):
                output.append(PreviewRow(number, keys, {"employee_id": employee_id}, "UNCHANGED", "SKIP", "HR_CORRECTION_PRESERVED", "Existing HR correction is preserved.")); continue
            parsed = parse_qualification_categories(authoritative_raw)
            record_payload = {"person_id": int(row["person_id"]), "employee_id": employee_id, "batch_id": batch_id, "row_id": row_id,
                              "entries": [item.__dict__ for item in parsed], "metadata_updated_at": str(row["metadata_updated_at"] or "")}
            fingerprint_material.append(record_payload)
            if not parsed:
                output.append(PreviewRow(number, keys, {"employee_id": employee_id}, "UNCHANGED", "SKIP", "NO_CATEGORY_IN_SOURCE", "No parsable category in source row.", record_payload)); continue
            output.append(PreviewRow(number, keys, {"employee_id": employee_id}, "CREATE", "UPSERT_CATEGORY_PROVENANCE", None, None, record_payload))
        return ScenarioResult(output, _fingerprint(fingerprint_material), [])

    def apply(self, conn, *, package_id: int, rows: list[PreviewRow], actor_user_id: int) -> dict[str, int]:
        changed = skipped = 0
        for row in rows:
            if row.result_group == "UNCHANGED": skipped += 1; continue
            if row.result_group != "CREATE" or not row.payload: raise WorkbookValidationError("Blocked preview rows cannot be applied.")
            payload = row.payload
            current = conn.execute(text("SELECT additional_profile FROM public.personnel_record_metadata WHERE person_id=:person_id FOR UPDATE"), {"person_id": payload["person_id"]}).scalar_one_or_none()
            profile = normalize_additional_profile(current or {})
            old = profile.get("qualification_categories") or []
            old = [entry for entry in old if not (isinstance(entry, dict) and (entry.get("provenance") or {}).get("source") == "data_exchange_control_list_category" and (entry.get("provenance") or {}).get("source_row_id") == payload["row_id"])]
            additions = []
            for entry in payload["entries"]:
                additions.append({"specialty": entry["specialty"], "category": entry["category"], "assigned_at": entry["assigned_at"],
                    "assigned_at_calculated": entry["assigned_at_calculated"], "review_status": entry["review_status"], "review_reason": entry["review_reason"],
                    "provenance": {"source": "data_exchange_control_list_category", "source_batch_id": payload["batch_id"], "source_row_id": payload["row_id"], "override_origin": "IMPORT", "source_fingerprint": entry["fingerprint"]}})
            profile["qualification_categories"] = old + additions
            conn.execute(text("""INSERT INTO public.personnel_record_metadata(person_id,additional_profile) VALUES(:person_id,CAST(:profile AS jsonb))
                ON CONFLICT(person_id) DO UPDATE SET additional_profile=EXCLUDED.additional_profile,updated_at=now()"""), {"person_id": payload["person_id"], "profile": json.dumps(profile, ensure_ascii=False)})
            changed += 1
        return {"created": changed, "updated": 0, "skipped": skipped}
