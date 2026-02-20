import os
import psycopg2
import openpyxl

# Получение строки подключения из переменной окружения или дефолтной
DATABASE_URL = os.getenv(
    "DATABASE_URL",
    "postgresql://postgres:postgres@localhost:5432/corpsite"
)

# Подключение к базе данных
conn = psycopg2.connect(DATABASE_URL)
cur = conn.cursor()

# Получение списка всех таблиц в схеме public
cur.execute("SELECT table_name FROM information_schema.tables WHERE table_schema = 'public'")
tables = [row[0] for row in cur.fetchall()]

wb = openpyxl.Workbook()
wb.remove(wb.active)  # удалить стандартный лист

for table in tables:
    cur.execute(f'SELECT * FROM "{table}"')
    rows = cur.fetchall()
    colnames = [desc[0] for desc in cur.description]
    ws = wb.create_sheet(title=table)
    ws.append(colnames)
    for row in rows:
        safe_row = []
        import datetime
        for value in row:
            # Если значение сложного типа — сериализуем в строку
            if isinstance(value, (dict, list, set)):
                safe_row.append(str(value))
            # Если это datetime с tzinfo — делаем naive или строку
            elif isinstance(value, datetime.datetime) and value.tzinfo is not None:
                # Можно убрать tzinfo (UTC), либо привести к строке
                safe_row.append(value.replace(tzinfo=None))
            elif isinstance(value, datetime.time) and value.tzinfo is not None:
                safe_row.append(value.replace(tzinfo=None))
            else:
                safe_row.append(value)
        ws.append(safe_row)

wb.save('all_tables_export.xlsx')
cur.close()
conn.close()

print(f'Экспортировано {len(tables)} таблиц в файл all_tables_export.xlsx')
