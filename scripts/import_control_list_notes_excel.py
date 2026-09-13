"""Safely import note facts from the locally supplied June control list.

The importer accepts a row only when its Excel IIN is unique, resolves exactly
one Person/Employee, exactly one row of the specified import batch, and the
Excel note equals that row's normalized ``note_raw``.  It never uses FIO.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

from app.db.engine import DATABASE_URL, engine
from app.services.hr_import_additional_status_service import parse_control_list_note, _fingerprint

NOTE_HEADER = "Примечание (декрет, инвалид, пенсионер)"
IIN_HEADER = "ИИН"
POLICY = "CONTROL_LIST_NOTE_STATUS_V1_EXCEL_VERIFIED"


def _loopback_only() -> None:
    if not any(host in DATABASE_URL.lower() for host in ("127.0.0.1", "localhost")):
        raise RuntimeError("refusing non-local database")


def import_notes(*, source: Path, batch_id: int) -> dict[str, int]:
    _loopback_only()
    wb = load_workbook(source, read_only=True, data_only=True)
    excel: list[tuple[str, str]] = []
    for ws in wb.worksheets:
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        columns = {str(value).strip(): index for index, value in enumerate(header) if value is not None}
        if IIN_HEADER not in columns or NOTE_HEADER not in columns:
            raise RuntimeError(f"required columns absent on sheet {ws.title!r}")
        for row in ws.iter_rows(min_row=2, values_only=True):
            iin, note = str(row[columns[IIN_HEADER]] or "").strip(), str(row[columns[NOTE_HEADER]] or "").strip()
            if iin.isdigit() and len(iin) == 12:
                excel.append((iin, note))
    counts = Counter(iin for iin, _ in excel)
    result = {"excel_rows": len(excel), "unmatched": 0, "mismatched_source": 0, "maternity_only": 0,
              "manual_review": 0, "created": 0, "already_present": 0}
    with engine.begin() as conn:
        people_by_iin: dict[str, list[dict]] = defaultdict(list)
        for row in conn.execute(text("""
            SELECT p.iin, p.person_id, e.employee_id
            FROM public.persons p JOIN public.employees e ON e.person_id=p.person_id
            WHERE p.iin IS NOT NULL
        """)).mappings():
            people_by_iin[str(row["iin"])].append(dict(row))
        batches: dict[str, list[dict]] = defaultdict(list)
        for row in conn.execute(text("""
            SELECT row_id, COALESCE(normalized_payload->>'iin', raw_payload->>'iin') AS iin,
                   normalized_payload->>'note_raw' AS note
            FROM public.hr_import_rows WHERE batch_id=:batch_id
        """), {"batch_id": batch_id}).mappings():
            batches[str(row["iin"] or "").strip()].append(dict(row))
        for iin, note in excel:
            if counts[iin] != 1 or len(people_by_iin[iin]) != 1 or len(batches[iin]) != 1:
                result["unmatched"] += 1
                continue
            source_row = batches[iin][0]
            if note != str(source_row["note"] or "").strip():
                result["mismatched_source"] += 1
                continue
            if not note:
                continue
            parsed = parse_control_list_note(note)
            if not parsed.facts:
                result["maternity_only" if not parsed.requires_manual_review else "manual_review"] += 1
                continue
            person = people_by_iin[iin][0]
            for fact in parsed.facts:
                exists = conn.execute(text("""
                    SELECT 1 FROM public.person_status_facts
                    WHERE source_row_id=:row_id AND fact_kind=:kind
                    LIMIT 1
                """), {"row_id": source_row["row_id"], "kind": fact.fact_kind}).scalar_one_or_none()
                if exists:
                    result["already_present"] += 1
                    continue
                conn.execute(text("""
                    INSERT INTO public.person_status_facts(
                      person_id, employee_context_id, fact_kind, effective_date, disability_group, icd10_code,
                      review_status, review_reason, source_batch_id, source_row_id, source_policy_version,
                      source_fingerprint, version)
                    VALUES(:person_id,:employee_id,:kind,:date,:group,:icd,:status,:reason,:batch,:row_id,:policy,:fingerprint,1)
                """), {"person_id": person["person_id"], "employee_id": person["employee_id"],
                          "kind": fact.fact_kind, "date": fact.effective_date, "group": fact.disability_group,
                          "icd": fact.icd10_code, "status": fact.review_status, "reason": fact.review_reason,
                          "batch": batch_id, "row_id": source_row["row_id"], "policy": POLICY,
                          "fingerprint": _fingerprint(note, fact)})
                result["created"] += 1
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--batch-id", type=int, required=True)
    args = parser.parse_args()
    print(import_notes(source=args.source, batch_id=args.batch_id))
