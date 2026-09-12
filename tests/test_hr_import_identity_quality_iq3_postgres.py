"""PostgreSQL integration coverage for IQ-3 import-time IIN classification."""
from __future__ import annotations

from pathlib import Path
from importlib.util import module_from_spec, spec_from_file_location
from contextlib import contextmanager
import os
from uuid import uuid4
import zipfile
from xml.etree import ElementTree

import pytest
from alembic.operations import Operations
from alembic.runtime.migration import MigrationContext
from openpyxl import Workbook
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url

from app.db.engine import engine as app_engine
from app.services.hr_import_control_list_storage import remove_stored_control_list_file
from app.services.hr_import_identity_quality_service import classify_import_identity_quality
from app.services.hr_import_service import import_control_list
from scripts.import_hr_control_list import SHEET_TYPE_RULES, parse_workbook


def _apply_iq2(conn, method: str) -> None:
    spec = spec_from_file_location("iq2_iq3_test", Path("alembic/versions/iq2a1b2c3d4_hr_import_identity_quality_foundation.py"))
    assert spec and spec.loader
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    with Operations.context(MigrationContext.configure(conn)):
        getattr(module, method)()


def _doctors_sheet_name() -> str:
    return SHEET_TYPE_RULES[2][1][0]


def _write_workbook(path: Path, rows: list[tuple[str, object]], *, date_iin: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _doctors_sheet_name()
    sheet.append(["", "department", "full_name", "", "iin", "", "", "", "", "position"])
    for offset, (full_name, iin) in enumerate(rows, start=2):
        sheet.cell(offset, 2, "IQ3")
        sheet.cell(offset, 3, full_name)
        cell = sheet.cell(offset, 5, iin)
        sheet.cell(offset, 10, "Doctor")
        if date_iin:
            cell.number_format = "yyyy-mm-dd"
    workbook.save(path)
    workbook.close()


def _write_formula_workbook_with_cached_iin(path: Path, *, full_name: str, cached_iin: str) -> None:
    """Write a formula cell whose OOXML cache must never be treated as an IIN."""
    _write_workbook(path, [(full_name, "=1")])
    with zipfile.ZipFile(path) as archive:
        members = {member.filename: archive.read(member.filename) for member in archive.infolist()}

    main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    sheet = ElementTree.fromstring(members["xl/worksheets/sheet1.xml"])
    cell = sheet.find(f".//{main}c[@r='E2']")
    assert cell is not None
    formula = cell.find(f"{main}f")
    assert formula is not None
    formula.text = "1"
    value = cell.find(f"{main}v")
    if value is None:
        value = ElementTree.SubElement(cell, f"{main}v")
    value.text = cached_iin
    cell.attrib.pop("t", None)
    members["xl/worksheets/sheet1.xml"] = ElementTree.tostring(
        sheet, encoding="utf-8", xml_declaration=True
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for filename, content in members.items():
            archive.writestr(filename, content)


@pytest.fixture
def import_transaction():
    source = make_url(os.environ["TEST_DATABASE_URL"])
    if source.drivername.split("+", 1)[0] != "postgresql" or (source.host or "").lower() not in {"127.0.0.1", "localhost", "::1"} or str(source.database) != "corpsite_test":
        raise RuntimeError("IQ-3 tests require corpsite_test only as a loopback disposable-clone template")
    target_name = f"corpsite_iq3_{uuid4().hex[:16]}_test"
    target = source.set(database=target_name)
    admin_engine = create_engine(source.set(database="postgres").render_as_string(hide_password=False), isolation_level="AUTOCOMMIT")
    target_engine = create_engine(target.render_as_string(hide_password=False))
    storage_refs: list[str] = []
    try:
        app_engine.dispose()
        with admin_engine.connect() as admin:
            admin.execute(text(f'CREATE DATABASE "{target_name}" TEMPLATE "{source.database}"'))
        connection = target_engine.connect()
        transaction = connection.begin()
        if connection.execute(text("SELECT to_regclass('public.hr_import_identity_review_events')")).scalar_one():
            connection.execute(text("DROP TABLE public.hr_import_identity_quality_states"))
            connection.execute(text("DROP TABLE public.hr_import_identity_review_events"))
            connection.execute(text("ALTER TABLE public.hr_import_rows DROP CONSTRAINT uq_hr_import_rows_batch_row"))
            connection.execute(text("ALTER TABLE public.employees DROP CONSTRAINT IF EXISTS uq_employees_employee_person"))
            connection.execute(text("DROP FUNCTION IF EXISTS public.prevent_hr_import_identity_review_event_mutation()"))
            connection.execute(text("DROP FUNCTION IF EXISTS public.assert_hr_import_iq_event_advances_state()"))
            connection.execute(text("DROP FUNCTION IF EXISTS public.assert_hr_import_iq_effective_event()"))
        _apply_iq2(connection, "upgrade")
        imported_by = int(connection.execute(text("SELECT user_id FROM users ORDER BY user_id LIMIT 1")).scalar_one())
        yield connection, storage_refs, imported_by
    finally:
        if "transaction" in locals():
            transaction.rollback()
        if "connection" in locals():
            connection.close()
        for storage_ref in storage_refs:
            remove_stored_control_list_file(storage_ref)
        target_engine.dispose()
        with admin_engine.connect() as admin:
            admin.execute(text("SELECT pg_terminate_backend(pid) FROM pg_stat_activity WHERE datname=:name AND pid<>pg_backend_pid()"), {"name": target_name})
            admin.execute(text(f'DROP DATABASE IF EXISTS "{target_name}"'))
        admin_engine.dispose()


def _import(conn, storage_refs: list[str], tmp_path: Path, imported_by: int, rows: list[tuple[str, object]], *, date_iin: bool = False) -> int:
    path = tmp_path / "контрольный2606.xlsx"
    _write_workbook(path, rows, date_iin=date_iin)
    batch_id, _, _ = import_control_list(conn, file_path=path, imported_by=imported_by)
    # IQ-2's effective-event constraints are deferred; exercise them through
    # the real pipeline before this fixture eventually rolls its transaction back.
    conn.execute(text("SET CONSTRAINTS ALL IMMEDIATE"))
    storage_ref = conn.execute(
        text("SELECT sf.storage_ref FROM hr_import_batches b JOIN hr_source_files sf ON sf.source_file_id=b.source_file_id WHERE b.batch_id=:batch"),
        {"batch": batch_id},
    ).scalar_one()
    storage_refs.append(str(storage_ref))
    return int(batch_id)


def _row(conn, batch_id: int, full_name: str) -> dict:
    return dict(conn.execute(
        text("SELECT row_id, employee_id, error_codes, normalized_payload FROM hr_import_rows WHERE batch_id=:batch AND normalized_payload->>'full_name'=:name"),
        {"batch": batch_id, "name": full_name},
    ).mappings().one())


@pytest.mark.parametrize(
    ("iin", "reason", "event_type"),
    [
        (None, "IIN_MISSING", "SYSTEM_IIN_MISSING"),
        ("123 456", "IIN_INVALID_FORMAT", "SYSTEM_IIN_INVALID_FORMAT"),
        ("123456789012", "IIN_UNMATCHED", "SYSTEM_IIN_UNMATCHED"),
    ],
)
def test_iq3_persists_system_identity_exception(import_transaction, tmp_path, iin, reason, event_type):
    conn, storage_refs, imported_by = import_transaction
    name = "Identity Exception"
    batch_id = _import(conn, storage_refs, tmp_path, imported_by, [(name, iin)])
    row = _row(conn, batch_id, name)
    state = conn.execute(text("SELECT current_event_id,identity_state,reason_code,person_id,employee_id FROM hr_import_identity_quality_states WHERE batch_id=:batch AND row_id=:row"), {"batch": batch_id, "row": row["row_id"]}).mappings().one()
    event = conn.execute(text("SELECT event_type,actor_type,actor_user_id,resulting_state,reason_code,person_id,employee_id FROM hr_import_identity_review_events WHERE event_id=:event"), {"event": state["current_event_id"]}).mappings().one()

    assert dict(state) | {"current_event_id": state["current_event_id"]} == {"current_event_id": state["current_event_id"], "identity_state": "UNRESOLVED", "reason_code": reason, "person_id": None, "employee_id": None}
    assert dict(event) == {"event_type": event_type, "actor_type": "SYSTEM", "actor_user_id": None, "resulting_state": "UNRESOLVED", "reason_code": reason, "person_id": None, "employee_id": None}
    assert row["error_codes"] is None
    assert conn.execute(text("SELECT error_rows,valid_rows,total_rows FROM hr_import_batches WHERE batch_id=:batch"), {"batch": batch_id}).one() == (0, 1, 1)


def _create_employee_with_iin(conn, *, full_name: str, iin: str) -> int:
    employee_id = int(conn.execute(text("INSERT INTO employees(full_name,is_active) VALUES(:name,true) RETURNING employee_id"), {"name": full_name}).scalar_one())
    conn.execute(text("INSERT INTO employee_identities(employee_id,identity_type,identity_value,is_primary) VALUES(:employee,'IIN',:iin,true)"), {"employee": employee_id, "iin": iin})
    return employee_id


def test_iq3_exact_iin_match_binds_without_identity_exception(import_transaction, tmp_path):
    conn, storage_refs, imported_by = import_transaction
    iin = "321654987012"
    employee_id = _create_employee_with_iin(conn, full_name="Canonical Employee", iin=iin)
    name = "Imported Employee"
    batch_id = _import(conn, storage_refs, tmp_path, imported_by, [(name, iin)])
    row = _row(conn, batch_id, name)
    assert row["employee_id"] == employee_id
    assert conn.execute(text("SELECT count(*) FROM hr_import_identity_quality_states WHERE batch_id=:batch"), {"batch": batch_id}).scalar_one() == 0
    assert conn.execute(text("SELECT count(*) FROM hr_import_identity_review_events WHERE batch_id=:batch"), {"batch": batch_id}).scalar_one() == 0


def test_iq3_never_uses_fio_fallback(import_transaction, tmp_path):
    conn, storage_refs, imported_by = import_transaction
    name = "Shared Name"
    _create_employee_with_iin(conn, full_name=name, iin="111111111111")
    batch_id = _import(conn, storage_refs, tmp_path, imported_by, [(name, "222222222222")])
    row = _row(conn, batch_id, name)
    assert row["employee_id"] is None
    assert conn.execute(text("SELECT reason_code FROM hr_import_identity_quality_states WHERE batch_id=:batch AND row_id=:row"), {"batch": batch_id, "row": row["row_id"]}).scalar_one() == "IIN_UNMATCHED"


def test_iq3_repeat_classification_is_idempotent(import_transaction, tmp_path):
    conn, storage_refs, imported_by = import_transaction
    name = "Repeat Person"
    batch_id = _import(conn, storage_refs, tmp_path, imported_by, [(name, None)])
    assert classify_import_identity_quality(conn, batch_id=batch_id) == 0
    assert conn.execute(text("SELECT count(*) FROM hr_import_identity_review_events WHERE batch_id=:batch"), {"batch": batch_id}).scalar_one() == 1
    assert conn.execute(text("SELECT count(*) FROM hr_import_identity_quality_states WHERE batch_id=:batch"), {"batch": batch_id}).scalar_one() == 1


def test_iq3_recovers_only_exact_raw_ooxml_iin(import_transaction, tmp_path):
    conn, storage_refs, imported_by = import_transaction
    name = "Ooxml Person"
    iin = "987654321098"
    path = tmp_path / "контрольный2606.xlsx"
    _write_workbook(path, [(name, int(iin))], date_iin=True)
    rows, _ = parse_workbook(path)
    assert rows[0].data["iin"] == iin and rows[0].iin_valid is True
    batch_id, _, _ = import_control_list(conn, file_path=path, imported_by=imported_by)
    conn.execute(text("SET CONSTRAINTS ALL IMMEDIATE"))
    storage_refs.append(str(conn.execute(text("SELECT sf.storage_ref FROM hr_import_batches b JOIN hr_source_files sf ON sf.source_file_id=b.source_file_id WHERE b.batch_id=:batch"), {"batch": batch_id}).scalar_one()))
    assert _row(conn, int(batch_id), name)["normalized_payload"]["iin"] == iin
    assert conn.execute(text("SELECT reason_code FROM hr_import_identity_quality_states WHERE batch_id=:batch"), {"batch": batch_id}).scalar_one() == "IIN_UNMATCHED"


def test_iq3_formula_cached_iin_is_invalid_and_never_auto_binds(import_transaction, tmp_path):
    conn, storage_refs, imported_by = import_transaction
    cached_iin = "987654321098"
    name = "Formula Person"
    _create_employee_with_iin(conn, full_name=name, iin=cached_iin)
    path = tmp_path / "контрольный2606.xlsx"
    _write_formula_workbook_with_cached_iin(path, full_name=name, cached_iin=cached_iin)

    parsed_rows, _ = parse_workbook(path)
    assert parsed_rows[0].data["iin"] == ""
    assert parsed_rows[0].iin_valid is False
    assert parsed_rows[0].iin_quality_issue == "IIN_INVALID_FORMAT"

    batch_id, _, _ = import_control_list(conn, file_path=path, imported_by=imported_by)
    conn.execute(text("SET CONSTRAINTS ALL IMMEDIATE"))
    storage_refs.append(str(conn.execute(text("SELECT sf.storage_ref FROM hr_import_batches b JOIN hr_source_files sf ON sf.source_file_id=b.source_file_id WHERE b.batch_id=:batch"), {"batch": batch_id}).scalar_one()))
    row = _row(conn, int(batch_id), name)
    assert row["employee_id"] is None
    assert row["normalized_payload"]["iin"] == ""
    assert conn.execute(text("SELECT reason_code FROM hr_import_identity_quality_states WHERE batch_id=:batch AND row_id=:row"), {"batch": batch_id, "row": row["row_id"]}).scalar_one() == "IIN_INVALID_FORMAT"


def test_iq3_identity_exceptions_do_not_change_structural_error_counts(import_transaction, tmp_path):
    conn, storage_refs, imported_by = import_transaction
    rows = [("Missing Person", None), ("Invalid Person", "12 34"), ("Unmatched Person", "444444444444")]
    batch_id = _import(conn, storage_refs, tmp_path, imported_by, rows)
    assert conn.execute(text("SELECT total_rows,valid_rows,error_rows FROM hr_import_batches WHERE batch_id=:batch"), {"batch": batch_id}).one() == (3, 3, 0)
    assert conn.execute(text("SELECT count(*) FROM hr_import_rows WHERE batch_id=:batch AND error_codes IS NOT NULL"), {"batch": batch_id}).scalar_one() == 0
    assert conn.execute(text("SELECT count(*) FROM hr_import_identity_quality_states WHERE batch_id=:batch"), {"batch": batch_id}).scalar_one() == 3
