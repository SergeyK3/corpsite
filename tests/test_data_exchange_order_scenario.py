from __future__ import annotations

import hashlib
import io

import pytest
from openpyxl import Workbook, load_workbook

from app.data_exchange.order_scenario import (
    ORDER_SCHEMA_VERSION,
    _REQUIRED_ORDER_COLUMNS,
    _open_xlsx,
    _read_order_rows,
    _safe_cell_text,
)


def _workbook_bytes(*, formula: bool = False, unsafe_path: bool = False) -> bytes:
    workbook = Workbook()
    orders = workbook.active
    orders.title = "Orders"
    orders.append(list(_REQUIRED_ORDER_COLUMNS))
    sha = hashlib.sha256(b"synthetic-pdf").hexdigest()
    values = {
        "schema_version": ORDER_SCHEMA_VERSION, "source_row_id": "synthetic-1",
        "employee_id": 101, "employee_full_name": "Synthetic Employee",
        "order_number_raw": "7-К", "order_number": "7-К", "order_date_raw": "2026-01-01",
        "order_date": "2026-01-01", "order_kind": "PERSONNEL", "action_type": "HIRE",
        "pdf_source": "../unsafe.pdf" if unsafe_path else "pdf/one.pdf", "pdf_sha256": sha,
        "word_source": "", "word_sha256": "", "source_text": "synthetic",
        "recognition_confidence": 0.99, "recognition_status": "READY",
        "matching_method": "EMPLOYEE_ID", "review_status": "PENDING", "note": "",
    }
    orders.append(["=1+1" if formula and key == "note" else values[key] for key in _REQUIRED_ORDER_COLUMNS])
    metadata = workbook.create_sheet("Metadata")
    metadata.append(["schema_version", ORDER_SCHEMA_VERSION])
    out = io.BytesIO(); workbook.save(out)
    return out.getvalue()


def test_order_workbook_accepts_versioned_synthetic_row() -> None:
    rows = _read_order_rows(_workbook_bytes())
    assert rows[0][0] == 2
    assert rows[0][1]["order_kind"] == "PERSONNEL"


def test_order_workbook_rejects_formula_and_zip_macro_shape() -> None:
    with pytest.raises(ValueError, match="formulas"):
        _read_order_rows(_workbook_bytes(formula=True))
    # A non-XLSX payload cannot be mistaken for an office workbook.
    with pytest.raises(ValueError, match="Unsupported"):
        _open_xlsx(b"not an xlsx")


def test_export_text_is_formula_safe() -> None:
    assert _safe_cell_text("=SUM(1,1)") == "'=SUM(1,1)"
    assert _safe_cell_text("ordinary") == "ordinary"


def test_workbook_keeps_source_path_for_scenario_validation() -> None:
    # Structural reader preserves evidence; path traversal is rejected later,
    # where the preview can associate an explicit row-level reason.
    rows = _read_order_rows(_workbook_bytes(unsafe_path=True))
    assert rows[0][1]["pdf_source"] == "../unsafe.pdf"
