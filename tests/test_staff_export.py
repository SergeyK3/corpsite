"""Database-free roster export contract checks (also runnable with unittest)."""
from io import BytesIO
import unittest
from unittest.mock import patch

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import get_current_user
from app.directory import employees_routes as routes
from app.services.staff_export import build_staff_workbook


class StaffExportTests(unittest.TestCase):
    def setUp(self):
        app = FastAPI()
        app.include_router(routes.router, prefix="/directory")
        app.dependency_overrides[get_current_user] = lambda: {"user_id": 1}
        self.client = TestClient(app)
        self.rows = [dict(
            id=str(i), fio=f"Иванов {i}", position={"name": "Врач"},
            department={"name": "Терапия"}, org_unit={"name": "Другое имя"},
            rate=0.5, status="active", group=2, unit=3, position_id=4, department_id=5,
        ) for i in range(1, 232)]
        self.rows.append(dict(self.rows[0], id="999", fio="Петров", group=8))
        self.calls = []

    def list_service(self, **args):
        self.calls.append(args)
        rows = self.rows
        for key, field in (("org_group_id", "group"), ("org_unit_id", "unit"), ("position_id", "position_id"), ("department_id", "department_id")):
            if args[key] is not None:
                rows = [row for row in rows if row[field] == args[key]]
        if args["q"]:
            rows = [row for row in rows if args["q"].lower() in row["fio"].lower() or args["q"] in row["id"]]
        if args["status"] != "all":
            rows = [row for row in rows if row["status"] == args["status"]]
        return {"items": rows[args["offset"]:args["offset"] + args["limit"]], "total": len(rows)}

    def test_export_matches_all_roster_pages(self):
        with patch.object(routes, "compute_scope", return_value={"scope_unit_id": 3, "scope_unit_ids": [3]}), patch.object(routes, "require_personnel_visibility_or_403"), patch.object(routes, "svc_list_employees", side_effect=self.list_service):
            for query, count in (
                ("status=all", 232),
                ("status=active&q=Иванов&org_group_id=2&org_unit_id=3&include_children=true&position_id=4&department_id=5", 231),
                ("status=all&q=231&org_group_id=2&position_id=4", 1),
            ):
                with self.subTest(query=query):
                    self.calls.clear()
                    # Pagination parameters must not truncate or shift the export.
                    response = self.client.get(f"/directory/employees/export.xlsx?{query}&offset=50&limit=50")
                    self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else "")
                    self.assertIn("spreadsheetml.sheet", response.headers["content-type"])
                    sheet = load_workbook(BytesIO(response.content)).active
                    values = list(sheet.values)
                    self.assertEqual(values[0], ("ФИО", "Должность", "Отделение", "Ставка", "Статус"))
                    self.assertEqual(len(values) - 1, count)
                    self.assertEqual(values[1][1:], ("Врач", "Терапия", 0.5, "Работает"))
                    self.assertEqual(self.calls[0]["offset"], 0)
                    self.assertEqual(self.calls[0]["scope_unit_ids"], [3])
                    self.assertFalse(self.calls[0]["include_applicants"])
                    if count > 200:
                        self.assertEqual(self.calls[1]["offset"], 200)
                    # Compare workbook names and order with the paginated roster.
                    roster = []
                    offset = 0
                    while True:
                        page = self.client.get(f"/directory/employees?{query}&limit=50&offset={offset}").json()
                        roster.extend(page["items"])
                        offset += 50
                        if offset >= page["total"]:
                            break
                    self.assertEqual([row[0] for row in values[1:]], [row["fio"] for row in roster])

    def test_empty_and_visibility(self):
        with patch.object(routes, "compute_scope", return_value={"scope_unit_id": None, "scope_unit_ids": []}), patch.object(routes, "require_personnel_visibility_or_403"), patch.object(routes, "svc_list_employees", return_value={"items": [], "total": 0}):
            response = self.client.get("/directory/employees/export.xlsx")
            self.assertEqual(response.status_code, 404)
            self.assertIn("Нет записей", response.json()["detail"])
        with patch.object(routes, "compute_scope", return_value={}), patch.object(routes, "require_personnel_visibility_or_403", side_effect=HTTPException(403)), patch.object(routes, "svc_list_employees") as service:
            self.assertEqual(self.client.get("/directory/employees/export.xlsx").status_code, 403)
            service.assert_not_called()

    def test_full_export_preserves_filters_and_scope_across_three_pages(self):
        self.rows.extend(dict(self.rows[0], id=str(i), fio=f"Иванов {i}") for i in range(300, 550))
        with patch.object(routes, "compute_scope", return_value={"scope_unit_id": 3, "scope_unit_ids": [3]}), patch.object(routes, "require_personnel_visibility_or_403") as visibility, patch.object(routes, "svc_list_employees", side_effect=self.list_service):
            response = self.client.get("/directory/employees/export.xlsx?status=active&q=Иванов&org_group_id=2&org_unit_id=3&include_children=true&position_id=4&department_id=5&limit=25&offset=100")
        self.assertEqual(response.status_code, 200)
        values = list(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual(len(values) - 1, 481)
        self.assertEqual([call["offset"] for call in self.calls], [0, 200, 400])
        self.assertEqual(visibility.call_count, 3)
        for call in self.calls:
            for key, value in dict(status="active", q="Иванов", org_group_id=2, org_unit_id=3, include_children=True, position_id=4, department_id=5, limit=200, scope_unit_id=3, scope_unit_ids=[3], include_applicants=False).items():
                self.assertEqual(call[key], value)
        self.assertEqual(values[1][0], "Иванов 1")
        self.assertEqual(values[-1][0], "Иванов 549")

    def test_export_requires_authentication(self):
        self.client.app.dependency_overrides.clear()
        with patch.object(routes, "svc_list_employees") as service:
            self.assertEqual(self.client.get("/directory/employees/export.xlsx").status_code, 401)
            service.assert_not_called()

    def test_literal_text_and_statuses(self):
        items = [
            {"fio": "=1+1", "status": "inactive", "date_to": None},
            {"status": "inactive", "date_to": "2026-01-01", "termination": {"verification_status": "UNVERIFIED"}},
            {"status": "archived"},
        ]
        sheet = load_workbook(BytesIO(build_staff_workbook(items))).active
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertEqual(sheet["A2"].value, "=1+1")
        self.assertEqual([sheet[f"E{i}"].value for i in (2, 3, 4)], ["Не работает", "Завершён; Сведения не верифицированы", "Архив"])
        self.assertEqual(sheet.auto_filter.ref, "A1:E4")
