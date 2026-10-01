from __future__ import annotations

import io
import pytest
from openpyxl import Workbook

from app.data_exchange.vacation_scenario import _load


def _book() -> bytes:
    book = Workbook(); book.remove(book.active)
    headers = ["Строка исходного Excel", "№ блока", "№ события", "Исходный текст", "Проверено кадровиком"]
    for title in ("Нет реквизитов", "Ошибки периодов", "Строка перевода"):
        sheet = book.create_sheet(title); sheet.append(["title"]); sheet.append(["instruction"]); sheet.append([]); sheet.append(headers)
        sheet.append(["42", 5, 1, "synthetic", "Нет"])
    classes = book.create_sheet("Классификация"); classes.append(["title"]); classes.append(["instruction"]); classes.append([]); classes.append(["Код категории"])
    output = io.BytesIO(); book.save(output); return output.getvalue()


def test_same_source_row_on_multiple_exception_sheets_merges_without_duplicate() -> None:
    records = _load(_book())
    assert len(records) == 1
    assert set(records[0]["sheets"]) == {"Нет реквизитов", "Ошибки периодов", "Строка перевода"}


def test_missing_required_sheet_blocks_input() -> None:
    book = Workbook(); book.active.title = "Нет реквизитов"; output = io.BytesIO(); book.save(output)
    with pytest.raises(ValueError, match="required worksheets"):
        _load(output.getvalue())
